"""PPTX slide body and notes extraction."""

from __future__ import annotations

from app.services.attachments.chunking import chunk_plain_sections
from app.services.attachments.models import ChunkLocation, DocumentSummary
from app.services.attachments.parsers import ParseResult


class PptxParser:
    def parse(self, data: bytes, *, attachment_id: str, filename: str) -> ParseResult:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError("缺少 PPTX 解析依赖 python-pptx，请安装 requirements.txt") from exc

        import io

        prs = Presentation(io.BytesIO(data))
        sections: list[tuple[str, ChunkLocation]] = []
        for slide_num, slide in enumerate(prs.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if not hasattr(shape, "text"):
                    continue
                text = (shape.text or "").strip()
                if text:
                    parts.append(text)
            notes = ""
            if slide.has_notes_slide and slide.notes_slide and slide.notes_slide.notes_text_frame:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            body = "\n".join(parts)
            if notes:
                body = f"{body}\n\n[备注]\n{notes}".strip()
            if body:
                sections.append((body, ChunkLocation(slide=slide_num)))
        chunks = chunk_plain_sections(attachment_id, sections)
        plain = "\n\n".join(text for text, _ in sections)
        return ParseResult(
            chunks=chunks,
            summary=DocumentSummary(
                slide_count=len(prs.slides),
                chunk_count=len(chunks),
                char_count=len(plain),
                searchable=len(chunks) > 0,
            ),
            plain_text=plain,
        )
