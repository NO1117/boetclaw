"""DOCX paragraph and table extraction — no macros or embedded objects."""

from __future__ import annotations

from app.services.attachments.chunking import chunk_plain_sections
from app.services.attachments.models import ChunkLocation, DocumentSummary
from app.services.attachments.parsers import ParseResult


class DocxParser:
    def parse(self, data: bytes, *, attachment_id: str, filename: str) -> ParseResult:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("缺少 DOCX 解析依赖 python-docx，请安装 requirements.txt") from exc

        import io

        doc = Document(io.BytesIO(data))
        sections: list[tuple[str, ChunkLocation]] = []
        para_count = 0
        for para in doc.paragraphs:
            text = (para.text or "").strip()
            if not text:
                continue
            para_count += 1
            style = (para.style.name if para.style else "") or f"段落 {para_count}"
            sections.append((text, ChunkLocation(section=style)))
        for table_idx, table in enumerate(doc.tables, start=1):
            rows: list[str] = []
            for row in table.rows:
                cells = [((cell.text or "").strip()) for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                sections.append(
                    (
                        "\n".join(rows),
                        ChunkLocation(section=f"表格 {table_idx}"),
                    )
                )
        chunks = chunk_plain_sections(attachment_id, sections)
        plain = "\n\n".join(text for text, _ in sections)
        return ParseResult(
            chunks=chunks,
            summary=DocumentSummary(
                paragraph_count=para_count,
                chunk_count=len(chunks),
                char_count=len(plain),
                searchable=len(chunks) > 0,
            ),
            plain_text=plain,
        )
